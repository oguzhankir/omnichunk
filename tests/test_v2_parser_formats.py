from __future__ import annotations

import io
from types import SimpleNamespace

import pytest

from omnichunk import Chunker
from omnichunk.context.entities import extract_entities
from omnichunk.context.rebase import rebase_context
from omnichunk.context.scope import build_scope_tree, find_scope_chain
from omnichunk.formats.docx_loader import load_docx_bytes
from omnichunk.formats.pdf import load_pdf_bytes
from omnichunk.parser import languages, tree_sitter
from omnichunk.types import ByteRange, ChunkContext, EntityInfo, EntityType, LineRange
from omnichunk.util.text_index import TextIndex


def test_docx_paragraph_heading_table_source_and_metadata() -> None:
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_heading("Türkçe heading", level=1)
    document.add_paragraph("")
    document.add_paragraph("你好 source paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "öğe"
    table.cell(1, 1).text = "42"
    document.add_table(rows=1, cols=1)
    stream = io.BytesIO()
    document.save(stream)
    loaded = load_docx_bytes(stream.getvalue())
    assert "你好 source paragraph." in loaded.text
    assert "Name\tValue\nöğe\t42" in loaded.text
    assert loaded.segments[0].metadata["heading_style"] == "Heading 1"
    assert loaded.segments[-1].metadata["docx_block"] == "table"
    assert not loaded.warnings


def test_docx_style_failure_preserves_text_and_warns(monkeypatch) -> None:
    docx = pytest.importorskip("docx")
    from docx.text.paragraph import Paragraph

    document = docx.Document()
    document.add_paragraph("Preserve this paragraph.")
    stream = io.BytesIO()
    document.save(stream)

    def broken_style(self):
        raise ValueError("style unavailable")

    monkeypatch.setattr(Paragraph, "style", property(broken_style))
    loaded = load_docx_bytes(stream.getvalue())
    assert loaded.text == "Preserve this paragraph.\n\n"
    assert loaded.warnings == ("paragraph_style_0",)


def test_pdf_legacy_extraction_and_failed_page_diagnostics(monkeypatch) -> None:
    from omnichunk.formats import pdf

    class LegacyPage:
        def extract_text(self):
            return "Türkçe extracted source."

    class FailedPage:
        def extract_text(self, **kwargs):
            raise ValueError("corrupt page")

    fake = SimpleNamespace(
        PdfReader=lambda stream: SimpleNamespace(pages=[LegacyPage(), FailedPage()])
    )
    monkeypatch.setattr(pdf.importlib, "import_module", lambda name: fake)
    loaded = load_pdf_bytes(b"fixture")
    assert loaded.text.startswith("Türkçe extracted source.")
    assert any("page_2" in warning for warning in loaded.warnings)


def test_pdf_missing_dependency_is_actionable(monkeypatch) -> None:
    from omnichunk.formats import pdf

    def absent(name):
        raise ImportError("unavailable")

    monkeypatch.setattr(pdf.importlib, "import_module", absent)
    with pytest.raises(ImportError, match="omnichunk\\[pdf\\]"):
        load_pdf_bytes(b"fixture")


def test_query_unavailable_uses_ast_mapping(monkeypatch) -> None:
    from omnichunk.context import entities

    source = "class Outer:\n    def inner(self):\n        return 1\n"
    parsed = tree_sitter.parse_code(source, "python")
    monkeypatch.setattr(entities, "_compile_query", lambda *args: None)
    found = extract_entities(source, "python", parsed.tree)
    assert {entity.name for entity in found} >= {"Outer", "inner"}
    assert next(entity for entity in found if entity.name == "inner").parent == "Outer"
    assert extract_entities("", "python", None) == []


def test_plugin_ast_reports_syntax_errors(monkeypatch) -> None:
    import omnichunk.plugins

    syntax_tree = SimpleNamespace(root_node=SimpleNamespace(has_error=True))
    monkeypatch.setattr(
        omnichunk.plugins, "get_parser", lambda language: lambda path, text: syntax_tree
    )
    parsed = tree_sitter.parse_code("bad source", "python")
    assert parsed.backend == "plugin"
    assert parsed.tree is syntax_tree
    assert parsed.errors == ["Tree contains syntax error nodes"]


def test_invalid_plugin_tree_recovers_with_builtin_parser(monkeypatch) -> None:
    import omnichunk.plugins

    monkeypatch.setattr(
        omnichunk.plugins, "get_parser", lambda language: lambda path, text: object()
    )
    parsed = tree_sitter.parse_code("x = 1\n", "python")
    assert parsed.backend == "tree-sitter"
    assert "no root node" in parsed.errors[0]


def test_tree_sitter_runtime_failure_has_deterministic_fallback(monkeypatch) -> None:
    class BrokenParser:
        def parse(self, raw):
            raise RuntimeError("instance-specific object detail")

    monkeypatch.setattr(tree_sitter, "get_ts_parser", lambda language: BrokenParser())
    source = "def oku():\n    return 1\n"
    first = tree_sitter.parse_code(source, "python")
    second = tree_sitter.parse_code(source, "python")
    assert first == second
    assert first.backend == "regex"
    assert first.errors == ["Tree-sitter parse failed: RuntimeError; fallback=regex"]


def test_parser_supports_constructor_requiring_language(monkeypatch) -> None:
    class ConstructorParser:
        def __init__(self, language):
            self.language = language

    marker = object()
    monkeypatch.setattr(languages, "_TSParser", ConstructorParser)
    assert languages._new_parser(marker).language is marker


def test_parser_unavailable_or_incompatible_returns_none(monkeypatch) -> None:
    monkeypatch.setattr(languages, "_TSParser", None)
    assert languages._new_parser(object()) is None

    class BrokenParser:
        def __init__(self, *args):
            raise RuntimeError("unsupported ABI")

    monkeypatch.setattr(languages, "_TSParser", BrokenParser)
    assert languages._new_parser(object()) is None


def test_deep_scope_lookup_and_line_only_context_rebasing() -> None:
    entities = [
        EntityInfo(f"scope{i}", EntityType.CLASS, byte_range=ByteRange(i, 5000 - i))
        for i in range(300)
    ]
    tree = build_scope_tree(entities)
    chain = find_scope_chain(tree, ByteRange(500, 501))
    assert len(chain) == 300
    assert chain[0].name == "scope299"
    context = ChunkContext(
        entities=[EntityInfo("external", EntityType.MODULE, line_range=LineRange(0, 1))]
    )
    index = TextIndex("Türkçe\n\nbody\n")
    rebased = rebase_context(context, index, len("Türkçe\n\n".encode()))
    assert rebased.entities[0].line_range == LineRange(2, 3)


def test_unclosed_fence_retains_code_without_false_heading() -> None:
    source = "# Heading\n\n```python\n# Code comment\nx = 1\n"
    chunks = Chunker(max_chunk_size=500, coverage_policy="lossless").chunk("source.md", source)
    assert "".join(chunk.text for chunk in chunks) == source
    assert any(chunk.context.language == "python" for chunk in chunks)
    assert all("Code comment" not in chunk.context.heading_hierarchy for chunk in chunks)


def test_public_loaded_document_enforces_budget_and_attaches_source() -> None:
    from omnichunk import ChunkOptions, FormatSegment, LoadedDocument, chunk_loaded_document

    source = "Türkçe canonical text. " * 15
    loaded = LoadedDocument(source, (FormatSegment(0, len(source), "prose"),), "fixture")
    options = ChunkOptions(
        max_chunk_size=32,
        min_chunk_size=1,
        size_unit="chars",
        context_mode="none",
        coverage_policy="lossless",
    )
    chunks = chunk_loaded_document("fixture.document", loaded, options)
    assert "".join(chunk.text for chunk in chunks) == source
    assert all(len(chunk.contextualized_text) <= 32 for chunk in chunks)
    assert all(
        chunk.source is not None and chunk.source.format_name == "fixture" for chunk in chunks
    )
    assert all(chunk.metadata["budget"]["overflow"] == 0 for chunk in chunks)
    assert chunks == chunk_loaded_document("fixture.document", loaded, options)
